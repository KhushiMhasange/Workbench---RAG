import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pytest
from fpdf import FPDF
from docx import Document

from rag.loaders import detect_type, load, sanitise_filename


def test_text_loader(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("hello\r\nworld\n", encoding="utf-8")
    kind, sections = load(p)
    assert kind == "txt" and len(sections) == 1
    assert "\r" not in sections[0].text


def test_markdown_loader(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# Title\nbody\n## Sub\nmore\n", encoding="utf-8")
    kind, sections = load(p)
    assert kind == "md" and len(sections) == 2
    assert sections[0].heading == "Title"


def test_pdf_loader(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="hello pdf page one")
    p = tmp_path / "a.pdf"
    pdf.output(str(p))
    kind, sections = load(p)
    assert kind == "pdf" and sections[0].page == 1
    assert "hello" in sections[0].text.lower()


def test_scanned_pdf_error(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    p = tmp_path / "blank.pdf"
    pdf.output(str(p))
    with pytest.raises(ValueError, match="no extractable text"):
        load(p)


def test_docx_loader(tmp_path):
    doc = Document()
    doc.add_heading("Intro", level=1)
    doc.add_paragraph("hello docx")
    tab = doc.add_table(rows=1, cols=2)
    tab.cell(0, 0).text = "a"
    tab.cell(0, 1).text = "b"
    p = tmp_path / "a.docx"
    doc.save(str(p))
    kind, sections = load(p)
    assert kind == "docx"
    assert any("hello" in s.text for s in sections)
    assert any("a | b" in s.text for s in sections)


def test_csv_loader(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("sku,price\nA1,10\nA2,20\n", encoding="utf-8")
    kind, sections = load(p)
    assert kind == "csv"
    assert "sku=A1" in sections[0].text


def test_sniff_wrong_extension(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="sniffed pdf")
    p = tmp_path / "fake.txt"
    pdf.output(str(p))
    assert detect_type(p) == "pdf"
    doc = Document()
    doc.add_paragraph("hi")
    q = tmp_path / "fake.txt"
    doc.save(str(q))
    assert detect_type(q) == "docx"


def test_unsupported(tmp_path):
    p = tmp_path / "a.bin"
    p.write_bytes(bytes(range(256)) * 4)
    with pytest.raises(ValueError, match="unsupported"):
        load(p)


def test_sanitise():
    assert sanitise_filename("../../etc/passwd") == "passwd"
    assert sanitise_filename("") == "file"
    assert sanitise_filename("...") == "file"
    assert sanitise_filename("héllo wörld?.txt") != ""
    assert len(sanitise_filename("a" * 200)) <= 100
    assert not sanitise_filename(".hidden").startswith(".")
