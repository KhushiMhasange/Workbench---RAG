from pathlib import Path

from docx import Document
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent / "fixtures"
ROOT.mkdir(parents=True, exist_ok=True)


def make_pdf():
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pages = [
        ("Welcome to Northwind Traders", "Northwind Traders employee handbook, revision 2026-03. All staff must read chapters 1 to 6."),
        ("Leave Policy", "Annual leave is 27 days per year. Sick leave requires ticket HR-4821. Parental leave is 16 weeks. Policy code LV-2026-A."),
        ("Expense Policy", "Meals up to $48 per day are reimbursed. Flights over $620 need approval ticket FIN-7710. Receipts required for all claims over $25."),
        ("Security Policy", "Passwords rotate every 90 days. VPN token SKU SEC-TOKEN-9X must be used offsite. Incident reports use ticket SEC-3301."),
        ("Office Hours", "Core hours are 10:00 to 15:00. Standup at 09:30. The Oslo office opens 08:00. Handbook version HB-2026-6."),
        ("Benefits", "Gym stipend is $75 per month. Lunch subsidy is $9 per day. Referral bonus is $1,500. Contact benefits@northwind.example."),
        ("Remote Work", "Remote allowance is $120 per month. Monitor SKU MON-27X4 is standard issue. Support ticket IT-9152 for setup."),
    ]
    for title, body in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(text=title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(w=0, text=body + " " + ("Additional guidance appears in the appendix. " * 8))
    pdf.output(str(ROOT / "handbook.pdf"))


def make_notes():
    text = """Kickoff 2026-02-04 -- attendees: Mara, Jonas, Priya.
Northwind launch date is 2026-09-14, budget $84,500. Decision: use ticket OPS-5510 for rollout.
Mara: vendor quote for monitor SKU MON-27X4 is $329, cheaper than $399 retail.
Random aside about cats to add noise, nothing relevant here.
Jonas: security audit due 2026-05-01, ticket SEC-3301 already filed (same as handbook, good).
Priya: pricing for Widget Pro (SKU WDG-1001) is $49, Widget Lite (SKU WDG-1002) is $29.
Long rambling paragraph about lunch orders and office plants that goes on and on without structure, just to make chunking work hard. """ * 4
    (ROOT / "notes.txt").write_text(text, encoding="utf-8")


def make_csv():
    rows = [
        "sku,name,price,stock",
        "WDG-1001,Widget Pro,49,120",
        "WDG-1002,Widget Lite,29,340",
        "MON-27X4,Standard Monitor,329,45",
        "SEC-TOKEN-9X,VPN Token,15,800",
        "GAD-5555,Gadget Plus,99,10",
    ]
    (ROOT / "pricing.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")


def make_docx():
    doc = Document()
    doc.add_heading("Travel Policy", level=1)
    doc.add_paragraph("Per-diem is $62 per day. Hotel cap is $180 per night. Ticket TRV-2044 for approvals.")
    doc.add_heading("Equipment", level=2)
    doc.add_paragraph("Standard laptop is SKU LAP-8810. Monitor SKU MON-27X4 is $329 (conflicts with notes retail claim, docx is authoritative).")
    tab = doc.add_table(rows=3, cols=3)
    data = [["SKU", "Item", "Price"], ["WDG-1001", "Widget Pro", "$49"], ["LAP-8810", "Laptop", "$1,250"]]
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            tab.cell(i, j).text = val
    doc.add_heading("Contacts", level=1)
    doc.add_paragraph("Travel desk: travel@northwind.example. Policy version DOC-2026-1.")
    doc.save(str(ROOT / "policy.docx"))


if __name__ == "__main__":
    make_pdf()
    make_notes()
    make_csv()
    make_docx()
    print("wrote fixtures to", ROOT)
