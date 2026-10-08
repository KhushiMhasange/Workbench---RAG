from pathlib import Path

from docx import Document

from ..models import Section


def load_docx(path: Path) -> list[Section]:
    doc = Document(str(path))
    sections: list[Section] = []
    buf: list[str] = []
    heading: str | None = None

    def flush():
        if buf and any(b.strip() for b in buf):
            sections.append(Section(text="\n".join(buf).strip(), heading=heading))

    for para in doc.paragraphs:
        style = ""
        try:
            style = para.style.name or ""
        except Exception:
            style = ""
        if style.startswith("Heading"):
            flush()
            buf = []
            heading = para.text.strip() or None
        elif para.text.strip():
            buf.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            line = " | ".join(c.text.strip() for c in row.cells).strip()
            if line.strip(" |"):
                buf.append(line)
    flush()
    return [s for s in sections if s.text.strip()]
